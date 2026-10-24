"""PPTX renderer: native editable shapes, multi-segment edges, determinism."""

from __future__ import annotations

import io
import zipfile

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from papercanvas.core.models import Figure
from papercanvas.layout import run_layout
from papercanvas.layout.routing import route_orthogonal
from papercanvas.render.pptx import render_pptx
from papercanvas.text.text_layout import TextLayout
from papercanvas.text.text_metrics import FontMetrics


@pytest.fixture
def laid_out(figure_dict: dict) -> Figure:
    figure_dict["nodes"] = [
        {"id": "a", "type": "rounded_rect", "label": "Input Video"},
        {"id": "b", "type": "rect", "label": "Video Encoder"},
        {"id": "c", "type": "rect", "label": "Prediction"},
    ]
    figure_dict["edges"] = [
        {"id": "e1", "source": "a", "target": "b"},
        {"id": "e2", "source": "b", "target": "c"},
    ]
    figure_dict["groups"] = [{"id": "g1", "children": ["a", "b"], "label": "Stage 1"}]
    return run_layout(Figure.model_validate(figure_dict), "linear-h")


def _render(fig: Figure) -> bytes:
    return render_pptx(fig, TextLayout(FontMetrics.for_family(fig.theme.font_family)))


def _presentation(fig: Figure) -> Presentation:
    return Presentation(io.BytesIO(_render(fig)))


def _shape_by_name(prs: Presentation, name: str):
    for shape in prs.slides[0].shapes:
        if shape.name == name:
            return shape
    return None


def test_pptx_is_byte_deterministic(laid_out: Figure) -> None:
    """Zip entries carry pinned timestamps -> identical bytes."""
    first = _render(laid_out.model_copy(deep=True))
    second = _render(laid_out.model_copy(deep=True))
    assert first == second


def test_pptx_contains_no_raster_media(laid_out: Figure) -> None:
    blob = _render(laid_out)
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        assert not [n for n in archive.namelist() if "media" in n.lower()]
    prs = _presentation(laid_out)
    for shape in prs.slides[0].shapes:
        assert shape.shape_type != MSO_SHAPE_TYPE.PICTURE


def test_pptx_uses_native_shapes_by_name(laid_out: Figure) -> None:
    prs = _presentation(laid_out)
    names = {shape.name for shape in prs.slides[0].shapes}
    assert {"background", "group:g1", "node:a", "node:b", "node:c"} <= names
    assert "edge:e1#seg0" in names and "edge:e2#seg0" in names


def test_pptx_node_geometry_matches_ir(laid_out: Figure) -> None:
    prs = _presentation(laid_out)
    for node in laid_out.nodes:
        shape = _shape_by_name(prs, f"node:{node.id}")
        assert shape is not None
        # EMU -> mm must land within Python float noise of the IR values
        assert shape.left / 36000 == pytest.approx(node.x, abs=1e-4)
        assert shape.top / 36000 == pytest.approx(node.y, abs=1e-4)
        assert shape.width / 36000 == pytest.approx(node.w, abs=1e-4)
        assert shape.height / 36000 == pytest.approx(node.h, abs=1e-4)


def test_pptx_edges_are_multi_segment_native_connectors(figure_dict: dict) -> None:
    """Constraint 4: one native straight connector per routed segment."""
    figure_dict["nodes"] = [
        {"id": "a", "x": 10, "y": 10, "w": 30, "h": 15},
        {"id": "b", "x": 80, "y": 40, "w": 30, "h": 15},
    ]
    figure_dict["edges"] = [{"id": "e1", "source": "a", "target": "b"}]
    fig = Figure.model_validate(figure_dict)
    prs = _presentation(fig)
    expected_points = route_orthogonal(fig.node("a"), "right", fig.node("b"), "left")
    expected_segments = len(expected_points) - 1

    connectors = [s for s in prs.slides[0].shapes if s.name.startswith("edge:e1#")]
    assert len(connectors) == expected_segments
    # each connector's endpoints match consecutive routed points (EMU -> mm)
    for index, connector in enumerate(connectors):
        start = expected_points[index]
        end = expected_points[index + 1]
        assert connector.begin_x / 36000 == pytest.approx(start.x, abs=1e-4)
        assert connector.begin_y / 36000 == pytest.approx(start.y, abs=1e-4)
        assert connector.end_x / 36000 == pytest.approx(end.x, abs=1e-4)
        assert connector.end_y / 36000 == pytest.approx(end.y, abs=1e-4)


def test_pptx_arrowhead_xml_present(figure_dict: dict) -> None:
    figure_dict["nodes"] = [
        {"id": "a", "x": 10, "y": 20, "w": 30, "h": 15},
        {"id": "b", "x": 80, "y": 20, "w": 30, "h": 15},
    ]
    figure_dict["edges"] = [{"id": "e1", "source": "a", "target": "b", "style": {"arrow": "end"}}]
    fig = Figure.model_validate(figure_dict)
    prs = _presentation(fig)
    connector = next(s for s in prs.slides[0].shapes if s.name.startswith("edge:e1#"))
    xml = connector.line._get_or_add_ln().xml
    assert 'type="triangle"' in xml and "tailEnd" in xml


def test_pptx_both_arrows_single_segment_respect_schema_order(figure_dict: dict) -> None:
    """Regression: headEnd must precede tailEnd in the a:ln sequence."""
    figure_dict["nodes"] = [
        {"id": "a", "x": 10, "y": 20, "w": 30, "h": 15},
        {"id": "b", "x": 80, "y": 20, "w": 30, "h": 15},
    ]
    figure_dict["edges"] = [{"id": "e1", "source": "a", "target": "b", "style": {"arrow": "both"}}]
    fig = Figure.model_validate(figure_dict)
    prs = _presentation(fig)
    connector = next(s for s in prs.slides[0].shapes if s.name.startswith("edge:e1#"))
    xml = connector.line._get_or_add_ln().xml
    assert xml.index("headEnd") < xml.index("tailEnd")


def test_pptx_group_style_none_matches_svg_defaulting(figure_dict: dict) -> None:
    """Regression: explicitly-nulled group style must resolve like SVG/PDF
    (#444444 solid), not keep group-default #999999 dashed."""
    figure_dict["nodes"][1].update({"x": 60.0, "y": 20.0, "w": 30.0, "h": 15.0})
    figure_dict["groups"] = [
        {"id": "g1", "children": ["a", "b"], "style": {"stroke": None, "stroke_dash": None}}
    ]
    fig = Figure.model_validate(figure_dict)
    prs = _presentation(fig)
    group = _shape_by_name(prs, "group:g1")
    xml = group.line._get_or_add_ln().xml
    assert 'val="444444"' in xml  # DEFAULT_STROKE, like SVG/PDF
    assert 'val="dash"' not in xml  # no dash: explicitly solid


def test_pptx_textboxes_carry_textlayout_lines(laid_out: Figure) -> None:
    prs = _presentation(laid_out)
    text = "\n".join(
        shape.text_frame.text
        for shape in prs.slides[0].shapes
        if shape.has_text_frame and shape.name == "text"
    )
    assert "Input Video" in text
    assert "Video Encoder" in text
    assert "Prediction" in text
    assert "Stage 1" in text


def test_pptx_slide_size_matches_canvas(laid_out: Figure) -> None:
    from papercanvas.render.pptx import SLIDE_MIN_MM

    prs = _presentation(laid_out)
    # slide size is the canvas, padded up to the OOXML 1-inch floor
    assert prs.slide_width / 36000 == pytest.approx(
        max(laid_out.canvas.width, SLIDE_MIN_MM), abs=1e-4
    )
    assert prs.slide_height / 36000 == pytest.approx(
        max(laid_out.canvas.height, SLIDE_MIN_MM), abs=1e-4
    )
    # node coordinates are unaffected by the padding (absolute EMU)
    node = laid_out.nodes[0]
    shape = _shape_by_name(prs, f"node:{node.id}")
    assert shape.top / 36000 == pytest.approx(node.y, abs=1e-4)


def test_pptx_core_properties_are_pinned(laid_out: Figure) -> None:
    import datetime

    prs = _presentation(laid_out)
    props = prs.core_properties
    assert props.created == datetime.datetime(2000, 1, 1, 0, 0, 0)
    assert props.modified == datetime.datetime(2000, 1, 1, 0, 0, 0)


def test_pptx_render_does_not_modify_source(laid_out: Figure, tmp_path) -> None:
    from papercanvas.core.document import FigureDocument

    path = tmp_path / "fig.json"
    FigureDocument.create(path, laid_out)
    before = path.read_bytes()
    _render(FigureDocument.load(path).figure)
    assert path.read_bytes() == before


def test_pptx_unplaced_nodes_refuse_to_render(figure_dict: dict) -> None:
    figure_dict["nodes"] = [{"id": "a", "label": "no coords"}]
    with pytest.raises(Exception) as excinfo:  # noqa: PT011 - RenderError expected
        _render(Figure.model_validate(figure_dict))
    assert "layout" in str(excinfo.value)
