"""Release Gate v0.1 — the eight acceptance items, automated where possible.

Items and where they are proven:
  1. IR validates legal and rejects illegal input      -> here + tests/test_schema.py
  2. one IR exports SVG, PDF and PPTX                  -> here
  3. SVG/PDF stay vector, PPTX stays editable          -> here
  4. all examples lint clean (no overflow/mislinks)    -> here
  5. recoloring one node keeps the others intact       -> here
  6. a failing patch changes neither file nor revision -> here
  7. agent completes create/modify/export without hand-editing JSON
                                                       -> tests/test_mcp_e2e.py
  8. CI green on Windows/macOS/Linux                   -> .github/workflows/ci.yml
                                                          (static check here,
                                                           live result on GitHub)
"""

from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from typer.testing import CliRunner

from papercanvas.cli import app

runner = CliRunner()

EXAMPLES = ["basic_pipeline.json", "video_unlearning.json", "multi_branch.json"]


def _copy_example(name: str, tmp_path: Path) -> Path:
    source = Path(__file__).resolve().parent.parent / "examples" / name
    target = tmp_path / name
    shutil.copy(source, target)
    return target


@pytest.mark.parametrize("example", EXAMPLES)
def test_gate_2_one_ir_exports_three_formats(example: str, tmp_path: Path) -> None:
    ir = _copy_example(example, tmp_path)
    result = runner.invoke(app, ["build", str(ir), "--format", "all"])
    assert result.exit_code == 0, result.output
    for suffix in (".svg", ".pdf", ".pptx"):
        artifact = ir.with_suffix(suffix)
        assert artifact.exists() and artifact.stat().st_size > 0


@pytest.mark.parametrize("example", EXAMPLES)
def test_gate_3_outputs_are_vector_and_editable(example: str, tmp_path: Path) -> None:
    ir = _copy_example(example, tmp_path)
    assert runner.invoke(app, ["build", str(ir), "--format", "all"]).exit_code == 0

    svg = ir.with_suffix(".svg").read_bytes()
    assert b"<image" not in svg
    root = etree.fromstring(svg)
    assert len(root.findall("{http://www.w3.org/2000/svg}rect")) >= 2

    pdf = ir.with_suffix(".pdf").read_bytes()
    assert pdf.startswith(b"%PDF")
    assert b"/Subtype /Image" not in pdf

    with zipfile.ZipFile(ir.with_suffix(".pptx")) as archive:
        assert not [n for n in archive.namelist() if "media" in n.lower()]
    from pptx import Presentation

    prs = Presentation(ir.with_suffix(".pptx"))
    shapes = list(prs.slides[0].shapes)
    assert any(s.has_text_frame for s in shapes)  # editable text
    assert any("edge:" in s.name for s in shapes)  # native edge connectors


@pytest.mark.parametrize("example", EXAMPLES)
def test_gate_4_examples_lint_clean(example: str, tmp_path: Path) -> None:
    ir = _copy_example(example, tmp_path)
    result = runner.invoke(app, ["lint", str(ir), "--format", "json"])
    assert result.exit_code == 0, result.output
    report = json.loads(result.output)
    assert report["summary"] == {"errors": 0, "warnings": 0}, report["findings"]


def test_gate_1_ir_validates_legal_and_rejects_illegal(tmp_path: Path) -> None:
    legal = _copy_example("basic_pipeline.json", tmp_path)
    assert runner.invoke(app, ["validate", str(legal)]).exit_code == 0

    payload = json.loads(legal.read_text(encoding="utf-8"))
    payload["edges"].append({"id": "e_bad", "source": "input", "target": "ghost"})
    illegal = tmp_path / "illegal.json"
    illegal.write_text(json.dumps(payload), encoding="utf-8")
    result = runner.invoke(app, ["validate", str(illegal)])
    assert result.exit_code == 1
    assert "invalid" in result.output


def test_gate_5_recoloring_one_node_keeps_the_others(tmp_path: Path) -> None:
    ir = _copy_example("basic_pipeline.json", tmp_path)
    before = json.loads(ir.read_text(encoding="utf-8"))

    ops = tmp_path / "recolor.json"
    ops.write_text(
        json.dumps(
            {"operations": [{"op": "set_style", "target": "input", "style": {"fill": "#FF0000"}}]}
        ),
        encoding="utf-8",
    )
    assert runner.invoke(app, ["patch", str(ir), "--ops", str(ops)]).exit_code == 0

    after = json.loads(ir.read_text(encoding="utf-8"))
    changed = []
    for old_node in before["nodes"]:
        new_node = next(n for n in after["nodes"] if n["id"] == old_node["id"])
        if new_node != old_node:
            changed.append(old_node["id"])
    assert changed == ["input"], f"unrequested changes to {changed}"


def test_gate_6_failed_patch_preserves_file_and_revision(tmp_path: Path) -> None:
    ir = _copy_example("basic_pipeline.json", tmp_path)
    before_bytes = ir.read_bytes()
    before_revision = json.loads(before_bytes.decode("utf-8"))["revision"]

    ops = tmp_path / "bad.json"
    ops.write_text(
        json.dumps({"operations": [{"op": "move_node", "id": "ghost", "x": 0, "y": 0}]}),
        encoding="utf-8",
    )
    result = runner.invoke(app, ["patch", str(ir), "--ops", str(ops)])
    assert result.exit_code == 1

    assert ir.read_bytes() == before_bytes
    assert json.loads(ir.read_text(encoding="utf-8"))["revision"] == before_revision


def test_gate_7_agent_flow_without_manual_json_editing(tmp_path: Path) -> None:
    """init -> patch -> layout -> render, driven only by CLI arguments."""
    ir = tmp_path / "agent_flow.json"
    assert runner.invoke(app, ["init", str(ir), "--id", "agent_demo"]).exit_code == 0

    ops = tmp_path / "ops.json"
    ops.write_text(
        json.dumps(
            {
                "operations": [
                    {"op": "add_node", "node": {"id": "x", "label": "X"}},
                    {"op": "add_node", "node": {"id": "y", "label": "Y"}},
                    {"op": "add_edge", "edge": {"id": "e", "source": "x", "target": "y"}},
                ]
            }
        ),
        encoding="utf-8",
    )
    assert runner.invoke(app, ["patch", str(ir), "--ops", str(ops)]).exit_code == 0
    assert runner.invoke(app, ["layout", str(ir), "--preset", "linear-h"]).exit_code == 0
    assert runner.invoke(app, ["render", str(ir), "--format", "all"]).exit_code == 0
    for suffix in (".svg", ".pdf", ".pptx"):
        assert ir.with_suffix(suffix).exists()


def test_gate_8_ci_matrix_covers_three_platforms() -> None:
    workflow = (
        Path(__file__).resolve().parent.parent / ".github" / "workflows" / "ci.yml"
    ).read_text(encoding="utf-8")
    for os_name in ("ubuntu-latest", "windows-latest", "macos-latest"):
        assert os_name in workflow
    for version in ("3.12", "3.14"):
        assert version in workflow
