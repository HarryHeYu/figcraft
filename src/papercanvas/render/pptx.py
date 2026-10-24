"""PPTX renderer (python-pptx) — native, fully editable shapes; never bitmaps.

Discipline:
- Constraint 4: an edge's orthogonal point list is rebuilt from native
  STRAIGHT connectors, one per segment — a single ELBOW connector can never
  reproduce an arbitrary routed path.
- Constraint 2: label lines come from TextLayout verbatim (one textbox per
  line); this renderer never wraps text.
- Units: mm -> EMU exclusively through core.geometry.
- Determinism: core properties are pinned and the zip container is rewritten
  with fixed entry timestamps, so the same IR yields byte-identical files.
"""

from __future__ import annotations

import datetime
import io
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from papercanvas.core.errors import RenderError
from papercanvas.core.geometry import mm_to_emu, pt_to_mm
from papercanvas.core.models import Figure
from papercanvas.layout.routing import route_orthogonal
from papercanvas.render.style import (
    DEFAULT_FILL,
    DEFAULT_TEXT_FILL,
    GROUP_LABEL_OFFSET_MM,
    GROUP_PADDING_MM,
    require_placed,
    resolve_fill,
    resolve_stroke,
    round_corner_mm,
    stroke_width_mm,
)
from papercanvas.text.text_layout import LINE_SPACING, TextLayout

_FIXED_TIMESTAMP = datetime.datetime(2000, 1, 1, 0, 0, 0)

#: OOXML floor: ST_SlideSizeCoordinate forbids slides smaller than 1 inch.
#: Smaller canvases keep their coordinates; the slide simply gets padding.
SLIDE_MIN_MM = 25.4

_DASH_STYLE = {
    "dashed": MSO_LINE.DASH,
    "dotted": MSO_LINE.ROUND_DOT,
    "solid": MSO_LINE.SOLID,
}


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    if len(h) == 3:
        h = "".join(ch * 2 for ch in h)
    return RGBColor.from_string(h.upper())


def _emu(mm: float) -> Emu:
    return Emu(mm_to_emu(mm))


def _apply_line_style(line, hex_color: str, width_mm: float, dash: str | None) -> None:
    line.color.rgb = _rgb(hex_color)
    line.width = _emu(width_mm)
    line.dash_style = _DASH_STYLE.get(dash or "solid", MSO_LINE.SOLID)


def _add_arrowhead(line, *, head: bool) -> None:
    """Arrowheads are outside python-pptx's API; set them on the a:ln element.

    CT_LineProperties is a schema sequence with headEnd BEFORE tailEnd, so a
    bare append would produce an invalid document on single-segment "both"
    arrows — insert in schema order instead.
    """
    ln = line._get_or_add_ln()
    tag = qn("a:headEnd") if head else qn("a:tailEnd")
    if ln.find(tag) is not None:
        return
    element = ln.makeelement(tag, {"type": "triangle", "w": "med", "len": "med"})
    if head:
        tail = ln.find(qn("a:tailEnd"))
        if tail is not None:
            tail.addprevious(element)
            return
    ln.append(element)


def _textbox_for_line(
    shapes, line, origin_x_mm: float, origin_y_mm: float,
    ascent_mm: float, line_height_mm: float, theme,
):
    box = shapes.add_textbox(
        _emu(origin_x_mm + line.x_mm),
        _emu(origin_y_mm + line.y_mm - ascent_mm),
        _emu(line.width_mm + 2.0),
        _emu(line_height_mm + 0.4),
    )
    box.name = "text"
    box.shadow.inherit = False
    frame = box.text_frame
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.word_wrap = False
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = line.text
    run.font.size = Pt(theme.font_size_pt)
    run.font.name = theme.font_family
    run.font.color.rgb = _rgb(DEFAULT_TEXT_FILL)
    return box


def _normalize_zip(blob: bytes) -> bytes:
    """Rewrite the pptx zip with fixed entry timestamps -> deterministic bytes."""
    source = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=_FIXED_TIMESTAMP.timetuple()[:6])
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = info.external_attr
            target.writestr(entry, source.read(info.filename))
    return out.getvalue()


def render_pptx(fig: Figure, text_layout: TextLayout) -> bytes:
    """Render the laid-out figure to a native-shape PPTX (bytes)."""
    require_placed(fig)
    theme = fig.theme
    ascent_pt, _ = text_layout.metrics.ascent_descent_pt(theme.font_size_pt)
    ascent_mm = pt_to_mm(ascent_pt)
    line_height_mm = pt_to_mm(theme.font_size_pt * LINE_SPACING)

    presentation = Presentation()
    presentation.slide_width = _emu(max(fig.canvas.width, SLIDE_MIN_MM))
    presentation.slide_height = _emu(max(fig.canvas.height, SLIDE_MIN_MM))
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes = slide.shapes

    props = presentation.core_properties
    props.author = "papercanvas"
    props.last_modified_by = "papercanvas"
    props.created = _FIXED_TIMESTAMP
    props.modified = _FIXED_TIMESTAMP
    props.revision = 1
    props.title = fig.id

    # background
    background = shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, _emu(fig.canvas.width), _emu(fig.canvas.height)
    )
    background.name = "background"
    background.shadow.inherit = False
    background.fill.solid()
    background.fill.fore_color.rgb = _rgb(resolve_fill(theme.background, theme, DEFAULT_FILL))
    background.line.fill.background()

    # groups first (background boxes), then edges, then nodes on top —
    # the same paint order as the SVG and PDF renderers
    for group in fig.groups:
        children = [fig.node(child_id) for child_id in group.children]
        missing = [
            c.id if c is not None else "<missing>"
            for c in children
            if c is None or not c.is_placed
        ]
        if missing:
            raise RenderError(
                f"group {group.id!r} contains unplaced or missing node(s) {missing}; "
                f"run `papercanvas layout` first"
            )
        x0 = min(c.x for c in children) - GROUP_PADDING_MM
        y0 = min(c.y for c in children) - GROUP_PADDING_MM
        x1 = max(c.x + c.w for c in children) + GROUP_PADDING_MM
        y1 = max(c.y + c.h for c in children) + GROUP_PADDING_MM
        box = shapes.add_shape(
            MSO_SHAPE.RECTANGLE, _emu(x0), _emu(y0), _emu(x1 - x0), _emu(y1 - y0)
        )
        box.name = f"group:{group.id}"
        box.shadow.inherit = False
        box.fill.background()
        # resolve exactly like SVG/PDF (None -> DEFAULT_STROKE, no dash) so
        # the three backends stay consistent even for explicitly-nulled styles
        _apply_line_style(
            box.line,
            resolve_stroke(group.style.stroke, theme),
            stroke_width_mm(group.style.stroke_width),
            group.style.stroke_dash,
        )
        if group.label:
            result = text_layout.layout_label(
                group.label,
                font_family=theme.font_family,
                font_size_pt=theme.font_size_pt,
                box_w_mm=x1 - x0,
                box_h_mm=GROUP_LABEL_OFFSET_MM,
                padding_h_mm=GROUP_PADDING_MM,
                padding_v_mm=0.5,
            )
            for line in result.lines:
                _textbox_for_line(
                    shapes, line, x0, y0, ascent_mm, line_height_mm, theme
                )

    for edge in fig.edges:
        source = fig.node(edge.source)
        target = fig.node(edge.target)
        if source is None or target is None:
            raise RenderError(
                f"edge {edge.id!r} references missing node(s) "
                f"{[n for n in (edge.source, edge.target) if fig.node(n) is None]}"
            )
        points = route_orthogonal(source, edge.source_port, target, edge.target_port)
        stroke = resolve_stroke(edge.style.stroke, theme)
        width_mm = stroke_width_mm(edge.style.stroke_width)
        for index, (start, end) in enumerate(zip(points, points[1:], strict=False)):
            connector = shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, _emu(start.x), _emu(start.y), _emu(end.x), _emu(end.y)
            )
            connector.name = f"edge:{edge.id}#seg{index}"
            connector.shadow.inherit = False
            _apply_line_style(connector.line, stroke, width_mm, None)
            if index == len(points) - 2 and edge.style.arrow in ("end", "both"):
                _add_arrowhead(connector.line, head=False)
            if index == 0 and edge.style.arrow in ("start", "both"):
                _add_arrowhead(connector.line, head=True)

    for node in fig.nodes:
        shape = shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE
            if node.type == "rounded_rect"
            else MSO_SHAPE.RECTANGLE,
            _emu(node.x),
            _emu(node.y),
            _emu(node.w),
            _emu(node.h),
        )
        shape.name = f"node:{node.id}"
        shape.shadow.inherit = False
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(resolve_fill(node.style.fill, theme, DEFAULT_FILL))
        _apply_line_style(
            shape.line,
            resolve_stroke(node.style.stroke, theme),
            stroke_width_mm(node.style.stroke_width),
            node.style.stroke_dash,
        )
        radius = round_corner_mm(node)
        if radius is not None:
            shape.adjustments[0] = min(0.5, radius / min(node.w, node.h))

        if node.label:
            result = text_layout.layout_label(
                node.label,
                font_family=theme.font_family,
                font_size_pt=theme.font_size_pt,
                box_w_mm=node.w,
                box_h_mm=node.h,
            )
            for line in result.lines:
                _textbox_for_line(
                    shapes, line, node.x, node.y, ascent_mm, line_height_mm, theme
                )

    buffer = io.BytesIO()
    presentation.save(buffer)
    return _normalize_zip(buffer.getvalue())
